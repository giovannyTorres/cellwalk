"""Measure XLSX loading with distant formatting and identical visible data."""

from __future__ import annotations

import argparse
import json
import platform
import tempfile
from pathlib import Path

import openpyxl
from benchmark import measure

from cellwalk import CellGrid


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--repeat", type=int, default=5)
    parser.add_argument("--output", type=Path, default=Path("benchmarks/xlsx_extent_results.json"))
    args = parser.parse_args()
    if args.repeat < 2:
        parser.error("repeat must be >=2")
    results: dict[str, object] = {
        "python": platform.python_version(),
        "platform": platform.platform(),
        "openpyxl": openpyxl.__version__,
        "repeats": args.repeat,
        "notes": "Identical 3x30 values. Timings include workbook parsing; memory is a separate "
        "tracemalloc execution, not process RSS. Declared rectangles are recorded without "
        "materializing their empty cells. No timing assertions are used in tests.",
    }
    with tempfile.TemporaryDirectory() as temporary:
        for name, styled in [
            ("clean", None),
            ("distant_row", (200_000, 30)),
            ("distant_column", (3, 16_384)),
        ]:
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            assert sheet is not None
            for row in range(1, 4):
                sheet.append([row * column for column in range(1, 31)])
            if styled is not None:
                sheet.cell(*styled).number_format = "0.00"
            declared_shape = [sheet.max_row, sheet.max_column]
            path = Path(temporary) / f"{name}.xlsx"
            workbook.save(path)
            workbook.close()
            grid = CellGrid.from_xlsx(path)
            assert (grid.nrows, grid.ncolumns) == (3, 30)
            assert all(
                grid.cell(row, column).value == (row + 1) * (column + 1)
                for row in range(3)
                for column in range(30)
            )
            results[name] = {
                "declared_shape": declared_shape,
                "declared_rectangle_cells": declared_shape[0] * declared_shape[1],
                "grid_cells": grid.nrows * grid.ncolumns,
                "from_xlsx": measure(lambda source=path: CellGrid.from_xlsx(source), args.repeat),
            }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(results, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(results, indent=2))


if __name__ == "__main__":
    main()
