"""Spatial cell navigation and explicit header association."""

from __future__ import annotations

from collections.abc import Mapping, Sequence
from dataclasses import dataclass, replace
from pathlib import Path
from typing import Any, Literal

import pandas as pd

Coordinate = tuple[int, int]
MissingPolicy = Literal["raise", "none"]


class GridError(ValueError):
    """Invalid coordinates, rules, or an unavailable header association."""


def _blank(value: Any) -> bool:
    return (isinstance(value, str) and not value.strip()) or (
        pd.api.types.is_scalar(value) and bool(pd.isna(value))
    )


def _integer(value: Any) -> bool:
    return isinstance(value, int) and not isinstance(value, bool)


@dataclass(frozen=True)
class Cell:
    """A physical cell and its original value anchor, all coordinates zero-based."""

    row: int
    column: int
    value: Any
    source_row: int
    source_column: int


@dataclass(frozen=True)
class CellGrid:
    """Immutable grid operations with explicit bounds, targets and associations."""

    _cells: tuple[Cell, ...]
    nrows: int
    ncolumns: int
    bounds: tuple[int, int, int, int]
    _targets: tuple[Cell, ...] = ()
    _headers: tuple[tuple[str, tuple[Cell | None, ...]], ...] = ()

    @classmethod
    def from_dataframe(cls, frame: pd.DataFrame) -> CellGrid:
        """Read values positionally; input index/column labels are not geometry."""
        if not isinstance(frame, pd.DataFrame):
            raise TypeError("frame must be a pandas DataFrame")
        nrows, ncolumns = frame.shape
        cells = tuple(
            Cell(row, column, frame.iat[row, column], row, column)
            for row in range(nrows)
            for column in range(ncolumns)
        )
        return cls(cells, nrows, ncolumns, (0, nrows, 0, ncolumns))

    @classmethod
    def from_xlsx(
        cls,
        path: str | Path,
        *,
        sheet: str | int = 0,
        merged: Literal["blank", "propagate"] = "blank",
    ) -> CellGrid:
        """Read cached XLSX values without evaluating formulas or preserving VBA.

        Merged children retain their physical coordinate and the anchor's source
        coordinate. Values remain blank unless propagation is explicitly enabled.
        Requires the optional ``xlsx`` extra (openpyxl).
        """
        if merged not in {"blank", "propagate"}:
            raise GridError("merged must be 'blank' or 'propagate'")
        try:
            from openpyxl import load_workbook  # type: ignore[import-untyped]
        except ImportError as exc:
            raise ImportError("Install cellwalk[xlsx] to read XLSX files") from exc
        workbook = load_workbook(path, data_only=True, keep_vba=False, keep_links=False)
        try:
            if (
                isinstance(sheet, bool)
                or not isinstance(sheet, (str, int))
                or isinstance(sheet, int)
                and not 0 <= sheet < len(workbook.worksheets)
                or isinstance(sheet, str)
                and sheet not in workbook.sheetnames
            ):
                raise GridError("sheet must be an existing worksheet name or zero-based index")
            worksheet = workbook.worksheets[sheet] if isinstance(sheet, int) else workbook[sheet]
            nrows, ncolumns = worksheet.max_row, worksheet.max_column
            anchors: dict[Coordinate, Coordinate] = {}
            for merged_range in worksheet.merged_cells.ranges:
                anchor = (merged_range.min_row - 1, merged_range.min_col - 1)
                for row in range(merged_range.min_row - 1, merged_range.max_row):
                    for column in range(merged_range.min_col - 1, merged_range.max_col):
                        anchors[(row, column)] = anchor
            cells = []
            for row in range(nrows):
                for column in range(ncolumns):
                    source_row, source_column = anchors.get((row, column), (row, column))
                    value = worksheet.cell(row + 1, column + 1).value
                    if merged == "propagate":
                        value = worksheet.cell(source_row + 1, source_column + 1).value
                    cells.append(Cell(row, column, value, source_row, source_column))
            return cls(tuple(cells), nrows, ncolumns, (0, nrows, 0, ncolumns))
        finally:
            workbook.close()

    def _inside(self, row: int, column: int) -> bool:
        top, bottom, left, right = self.bounds
        return top <= row < bottom and left <= column < right

    def cell(self, row: int, column: int) -> Cell:
        """Return a cell inside the active navigation rectangle."""
        if not _integer(row) or not _integer(column) or not self._inside(row, column):
            raise GridError(f"Coordinate {(row, column)} lies outside grid bounds {self.bounds}")
        return self._cells[row * self.ncolumns + column]

    def within(
        self, row_start: int, row_stop: int, column_start: int, column_stop: int
    ) -> CellGrid:
        """Restrict navigation to a half-open subrectangle; clear data selection."""
        values = (row_start, row_stop, column_start, column_stop)
        top, bottom, left, right = self.bounds
        if (
            not all(_integer(value) for value in values)
            or not top <= row_start <= row_stop <= bottom
            or not left <= column_start <= column_stop <= right
        ):
            raise GridError(
                "within bounds must form a half-open subrectangle of the current bounds"
            )
        return replace(self, bounds=values, _targets=(), _headers=())

    def data(
        self,
        *,
        rows: Sequence[int],
        columns: Sequence[int],
        include_blank: bool = False,
    ) -> CellGrid:
        """Select the Cartesian product of explicit row/column positions.

        Target order follows supplied rows, then supplied columns. Reselection
        clears associations. Structural blank targets are omitted by default.
        """
        selected_rows, selected_columns = list(rows), list(columns)
        top, bottom, left, right = self.bounds
        for positions, start, stop in (
            (selected_rows, top, bottom),
            (selected_columns, left, right),
        ):
            if any(not _integer(pos) or not start <= pos < stop for pos in positions) or len(
                set(positions)
            ) != len(positions):
                raise GridError("data positions must be distinct integer coordinates inside bounds")
        targets = tuple(
            self.cell(row, column)
            for row in selected_rows
            for column in selected_columns
            if include_blank or not _blank(self.cell(row, column).value)
        )
        return replace(self, _targets=targets, _headers=())

    def _walk(self, row: int, column: int, delta: Coordinate, skip_blank: bool) -> Cell | None:
        self.cell(row, column)
        row, column = row + delta[0], column + delta[1]
        while self._inside(row, column):
            candidate = self.cell(row, column)
            if not skip_blank or not _blank(candidate.value):
                return candidate
            row, column = row + delta[0], column + delta[1]
        return None

    def above(self, row: int, column: int, *, skip_blank: bool = True) -> Cell | None:
        """Find the nearest cell above, optionally skipping structural blanks."""
        return self._walk(row, column, (-1, 0), skip_blank)

    def below(self, row: int, column: int, *, skip_blank: bool = True) -> Cell | None:
        """Find the nearest cell below, optionally skipping structural blanks."""
        return self._walk(row, column, (1, 0), skip_blank)

    def left(self, row: int, column: int, *, skip_blank: bool = True) -> Cell | None:
        """Find the nearest cell to the left, optionally skipping blanks."""
        return self._walk(row, column, (0, -1), skip_blank)

    def right(self, row: int, column: int, *, skip_blank: bool = True) -> Cell | None:
        """Find the nearest cell to the right, optionally skipping blanks."""
        return self._walk(row, column, (0, 1), skip_blank)

    def _attach(
        self, name: str, headers: tuple[Cell | None, ...], missing: MissingPolicy
    ) -> CellGrid:
        if missing not in {"raise", "none"}:
            raise GridError("missing must be 'raise' or 'none'")
        if not isinstance(name, str) or not name:
            raise GridError("Header name must be a nonempty string")
        existing = [old_name for old_name, _ in self._headers]
        names = ["row", "column", "source_row", "source_column", "value"]
        for label in [*existing, name]:
            names.extend(
                [
                    label,
                    f"{label}_row",
                    f"{label}_column",
                    f"{label}_source_row",
                    f"{label}_source_column",
                ]
            )
        if len(set(names)) != len(names):
            raise GridError("Header name collides with an existing value or coordinate column")
        if missing == "raise" and any(header is None or _blank(header.value) for header in headers):
            raise GridError(
                f"Missing header {name!r}; choose missing='none' to retain unassigned targets"
            )
        normalized = tuple(
            None if header is None or _blank(header.value) else header for header in headers
        )
        return replace(self, _headers=(*self._headers, (name, normalized)))

    def headers_above(
        self,
        name: str,
        *,
        row: int,
        span: Literal["none", "right"] = "none",
        missing: MissingPolicy = "raise",
    ) -> CellGrid:
        """Associate a fixed header row, optionally spanning blank cells rightward."""
        top, bottom, left, _ = self.bounds
        if not _integer(row) or not top <= row < bottom or span not in {"none", "right"}:
            raise GridError("Specify a header row inside bounds and span='none' or 'right'")
        headers: list[Cell | None] = []
        for target in self._targets:
            if row >= target.row:
                raise GridError("headers_above requires the header row above every target")
            candidate = self.cell(row, target.column)
            header: Cell | None = candidate
            if _blank(candidate.value) and span == "right":
                header = next(
                    (
                        self.cell(row, col)
                        for col in range(target.column - 1, left - 1, -1)
                        if not _blank(self.cell(row, col).value)
                    ),
                    None,
                )
            headers.append(header)
        return self._attach(name, tuple(headers), missing)

    def headers_left(
        self,
        name: str,
        *,
        column: int,
        span: Literal["none", "down"] = "none",
        missing: MissingPolicy = "raise",
    ) -> CellGrid:
        """Associate a fixed header column, optionally spanning blanks downward."""
        top, _, left, right = self.bounds
        if not _integer(column) or not left <= column < right or span not in {"none", "down"}:
            raise GridError("Specify a header column inside bounds and span='none' or 'down'")
        headers: list[Cell | None] = []
        for target in self._targets:
            if column >= target.column:
                raise GridError("headers_left requires the header column left of every target")
            candidate = self.cell(target.row, column)
            header: Cell | None = candidate
            if _blank(candidate.value) and span == "down":
                header = next(
                    (
                        self.cell(r, column)
                        for r in range(target.row - 1, top - 1, -1)
                        if not _blank(self.cell(r, column).value)
                    ),
                    None,
                )
            headers.append(header)
        return self._attach(name, tuple(headers), missing)

    def associate(
        self,
        name: str,
        mapping: Mapping[Coordinate, Coordinate],
        *,
        missing: MissingPolicy = "raise",
    ) -> CellGrid:
        """Attach arbitrary explicitly mapped header cells to selected targets."""
        targets = {(target.row, target.column) for target in self._targets}
        if set(mapping) - targets:
            raise GridError("Mapping contains a coordinate outside the selected data cells")
        if any(
            not isinstance(coordinate, tuple) or len(coordinate) != 2
            for coordinate in mapping.values()
        ):
            raise GridError("Mapped header coordinates must be (row, column) tuples")
        headers = tuple(
            self.cell(*mapping[(target.row, target.column)])
            if (target.row, target.column) in mapping
            else None
            for target in self._targets
        )
        return self._attach(name, headers, missing)

    def to_frame(self) -> pd.DataFrame:
        """Return values and physical/anchor provenance without dtype inference."""
        columns = ["row", "column", "source_row", "source_column", "value"]
        for name, _ in self._headers:
            columns.extend(
                [
                    name,
                    f"{name}_row",
                    f"{name}_column",
                    f"{name}_source_row",
                    f"{name}_source_column",
                ]
            )
        records = []
        for pos, target in enumerate(self._targets):
            record = [
                target.row,
                target.column,
                target.source_row,
                target.source_column,
                target.value,
            ]
            for _, headers in self._headers:
                header = headers[pos]
                record.extend(
                    [
                        header.value,
                        header.row,
                        header.column,
                        header.source_row,
                        header.source_column,
                    ]
                    if header is not None
                    else [None] * 5
                )
            records.append(record)
        return pd.DataFrame(records, columns=columns, dtype=object)
